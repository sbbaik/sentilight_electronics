# -*- coding: utf-8 -*-
"""Build the fill-in workbook for the audited request set (Reviewer 3, point 1).

One row per executable policy cell (240), shuffled so the writer does not produce thirty red
sentences in a row. The gold tuple is fixed by the cell, so the writer has nothing to label:
only the request sentence is entered. Variation types answer the reviewer's six axes —
colour, emotion and intensity are the cell factors; omission is carried by the "none" cells;
negation and word order are assigned as writing instructions, with a context-clause type
added because it is the shape most real requests take.

Negation is only assigned where the negated intensity has one reading (dark ↔ "not bright",
bright ↔ "not dark", soft ↔ "not strong", strong ↔ "not weak"); it is never assigned to the
neutral intensity, where "not bright" could mean soft or dark and the gold would be ambiguous.
"""
from __future__ import annotations

import json
import random
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BACKEND = Path.home() / "Sentilight_Paper_2/Implement_New_Android_v6/backend"
sys.path.insert(0, str(BACKEND / "models/compact_llm"))
sys.path.insert(0, str(BACKEND))
import training_data as td  # noqa: E402

HERE = Path(__file__).resolve().parents[1]
OUT_XLSX = HERE / "data/audited_requests/audit_sheet_K240.xlsx"
OUT_GOLD = HERE / "data/audited_requests/audited_cells_gold.json"

SEED = 20261021
DIMS = ("H", "S", "B", "Dimmer", "CT")

COLOUR_KO = {k: v["label"] for k, v in td.COLOR_PROFILES.items()}
EMOTION_KO = {"sadness": "슬픔·우울", "happiness": "행복·기쁨", "anger": "화남·짜증",
              "calm": "평온·편안", "anxiety": "불안·긴장"}
INTENSITY_KO = {"neutral": "지정 없음", "soft": "은은하게", "strong": "강하게",
                "bright": "밝게", "dark": "어둡게"}
NEGATION_HINT = {"dark": "'밝지 않게' 같은 부정형으로", "bright": "'어둡지 않게' 같은 부정형으로",
                 "soft": "'너무 세지 않게' 같은 부정형으로", "strong": "'약하지 않게' 같은 부정형으로"}

VARIATION = {
    "일반": "평소 스마트 스피커에 말하듯 자연스럽게",
    "부정형": None,                       # filled per intensity
    "어순변형": "조명 요청을 문장 맨 앞이나 중간에 두고, 나머지를 뒤에",
    "맥락추가": "상황을 한 구절 덧붙이기 (예: '회의 끝나고 집에 왔는데')",
}
QUOTA = {"부정형": 40, "어순변형": 40, "맥락추가": 40}   # the rest are 일반

FONT = "Arial"
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="EEEEEE")
HEAD = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def enumerate_cells():
    cells = []
    for colour in [None] + list(td.COLOR_PROFILES):
        for emotion in [None] + list(td.EMOTION_PROFILES):
            for intensity in td.INTENSITY_MODIFIERS:
                out = td.normalize_output(td.resolve_case_output(
                    {"base_color": colour, "emotion": emotion, "intensity": intensity}))
                cells.append({"colour": colour, "emotion": emotion, "intensity": intensity,
                              "gold": {d: int(out[d]) for d in DIMS}})
    assert len(cells) == 240
    for i, c in enumerate(cells, 1):
        c["cell_id"] = f"K240-{i:03d}"
    return cells


def assign_variations(cells, rng):
    """Spread each variation type across colours, emotions and intensities."""
    pool = list(range(len(cells)))
    rng.shuffle(pool)
    for c in cells:
        c["variation"] = "일반"
    # negation: only non-neutral intensity, drawn evenly across the four intensities
    neg_pool = {k: [i for i in pool if cells[i]["intensity"] == k] for k in NEGATION_HINT}
    per = QUOTA["부정형"] // len(NEGATION_HINT)
    for k, ids in neg_pool.items():
        for i in ids[:per]:
            cells[i]["variation"] = "부정형"
    rest = [i for i in pool if cells[i]["variation"] == "일반"]
    for i in rest[:QUOTA["어순변형"]]:
        cells[i]["variation"] = "어순변형"
    for i in rest[QUOTA["어순변형"]:QUOTA["어순변형"] + QUOTA["맥락추가"]]:
        cells[i]["variation"] = "맥락추가"


def hint_for(c):
    if c["variation"] == "부정형":
        return NEGATION_HINT[c["intensity"]]
    return VARIATION[c["variation"]]


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True)
        cell.fill = HEAD
        cell.border = BOX
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def build():
    rng = random.Random(SEED)
    cells = enumerate_cells()
    assign_variations(cells, rng)
    order = list(range(len(cells)))
    rng.shuffle(order)

    wb = Workbook()

    # ------------------------------------------------------------------ 작성안내
    ws = wb.active
    ws.title = "작성안내"
    lines = [
        ("사람 작성 평가 문장 — 작성 안내", True),
        ("", False),
        ("목적: 학습 템플릿과 무관하게 사람이 쓴 한국어 조명 요청 240문장으로 모델을 평가합니다. "
         "심사자(Reviewer 3)가 색상·감정·강도·부정·어순·생략별 층화 보고를 요구했습니다.", False),
        ("", False),
        ("채울 곳: '문장작성' 시트의 노란색 H열(요청 문장)만 채우시면 됩니다. 다른 열은 수정하지 마십시오.", True),
        ("", False),
        ("규칙", True),
        ("1. 각 행의 색·감정·강도가 그 문장의 정답입니다. 문장을 읽는 사람이 그 조합을 알아들을 수 있어야 합니다.", False),
        ("2. 색이 '없음'이면 색을 말하지 마십시오. 감정이 '없음'이면 감정을 말하지 마십시오. "
         "둘 다 없으면 강도만 말하는 요청입니다.", False),
        ("3. 강도가 '지정 없음'이면 밝기·세기를 말하지 마십시오.", False),
        ("4. '변형 유형'과 '작성 힌트'를 따르되, 그 셀에 어색하면 자연스러운 쪽을 택하고 I열(메모)에 적어 주십시오.", False),
        ("5. 학습 템플릿·학습 데이터를 보지 말고 쓰십시오. 평소 말투 그대로가 가장 좋습니다. 길이는 자유입니다.", False),
        ("6. 같은 표현을 반복하지 마십시오. 행이 섞여 있는 것은 그 때문입니다.", False),
        ("", False),
        ("작성 예시 (실제 행 아님)", True),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=12 if r == 1 else 10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ex_head = ["번호", "셀ID", "색", "감정", "강도", "변형 유형", "작성 힌트", "요청 문장 (예시)", "메모"]
    ex_row = ["예", "K240-000", "주황색", "슬픔·우울", "은은하게", "맥락추가",
              "상황을 한 구절 덧붙이기", "회의 끝나고 집에 왔는데 우울해. 조명은 은은한 주황빛으로 해줄래", ""]
    r0 = len(lines) + 1
    for col, v in enumerate(ex_head, 1):
        ws.cell(row=r0, column=col, value=v)
    style_header(ws, r0, len(ex_head))
    for col, v in enumerate(ex_row, 1):
        c = ws.cell(row=r0 + 1, column=col, value=v)
        c.font = Font(name=FONT)
        c.border = BOX
        if col == 8:
            c.fill = YELLOW
    ws.column_dimensions["A"].width = 110
    for col in "BCDEFGHI":
        ws.column_dimensions[col].width = 14
    ws.column_dimensions["H"].width = 60

    # ------------------------------------------------------------------ 문장작성
    ws = wb.create_sheet("문장작성")
    head = ["번호", "셀ID", "색", "감정", "강도", "변형 유형", "작성 힌트", "요청 문장", "메모"]
    ws.append(head)
    style_header(ws, 1, len(head))
    for n, idx in enumerate(order, 1):
        c = cells[idx]
        c["row_no"] = n
        ws.append([
            n, c["cell_id"],
            COLOUR_KO.get(c["colour"], "없음"),
            EMOTION_KO.get(c["emotion"], "없음"),
            INTENSITY_KO[c["intensity"]],
            c["variation"], hint_for(c), "", "",
        ])
        r = n + 1
        for col in range(1, len(head) + 1):
            cell = ws.cell(row=r, column=col)
            cell.font = Font(name=FONT)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=(col in (7, 8, 9)))
            if col == 8:
                cell.fill = YELLOW
            elif col <= 7:
                cell.fill = GREY
    widths = {"A": 6, "B": 11, "C": 9, "D": 11, "E": 11, "F": 10, "G": 34, "H": 62, "I": 24}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "H2"
    ws.cell(row=1, column=8).comment = Comment(
        "이 열만 채워 주십시오. 각 행의 색·감정·강도가 정답입니다.", "SentiLight")
    dv = DataValidation(type="list", formula1='"일반,부정형,어순변형,맥락추가"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{len(cells) + 1}")
    # a visible count of how many rows are still empty, so the writer can see progress
    ws.cell(row=len(cells) + 3, column=7, value="남은 행 수").font = Font(name=FONT, bold=True)
    ws.cell(row=len(cells) + 3, column=8,
            value=f'=COUNTBLANK(H2:H{len(cells) + 1})').font = Font(name=FONT, bold=True)

    # ------------------------------------------------------------------ 범주표
    ws = wb.create_sheet("범주표")
    ws.append(["구분", "시트 표기", "정책 내부 키", "뜻 / 예시 표현"])
    style_header(ws, 1, 4)
    for k, ko in COLOUR_KO.items():
        ws.append(["색", ko, k, ""])
    ws.append(["색", "없음", "none", "색을 말하지 않음"])
    for k, ko in EMOTION_KO.items():
        ws.append(["감정", ko, k, ", ".join(td.EMOTION_PROFILES[k]["phrases"][:3])])
    ws.append(["감정", "없음", "none", "감정을 말하지 않음"])
    for k, ko in INTENSITY_KO.items():
        ws.append(["강도", ko, k, td.INTENSITY_MODIFIERS[k]["label"] or "밝기·세기를 말하지 않음"])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
            cell.border = BOX
    for col, w in zip("ABCD", (8, 14, 14, 40)):
        ws.column_dimensions[col].width = w

    # ------------------------------------------------------------------ gold (수정 금지)
    ws = wb.create_sheet("gold_수정금지")
    ws.append(["번호", "셀ID", "colour", "emotion", "intensity", "variation"] + list(DIMS))
    style_header(ws, 1, 6 + len(DIMS))
    for idx in order:
        c = cells[idx]
        ws.append([c["row_no"], c["cell_id"], c["colour"] or "none", c["emotion"] or "none",
                   c["intensity"], c["variation"]] + [c["gold"][d] for d in DIMS])
    ws.cell(row=1, column=1).comment = Comment(
        "정책이 계산한 정답 튜플입니다. 분석 단계에서 씁니다. 수정하지 마십시오.", "SentiLight")
    ws.protection.sheet = True
    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 11

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    OUT_GOLD.write_text(json.dumps(
        {"seed": SEED, "n": len(cells),
         "variation_counts": {v: sum(1 for c in cells if c["variation"] == v) for v in VARIATION},
         "rows": sorted(cells, key=lambda c: c["row_no"])},
        indent=1, ensure_ascii=False), encoding="utf-8")
    counts = {v: sum(1 for c in cells if c["variation"] == v) for v in VARIATION}
    print(f"  {OUT_XLSX.name}: {len(cells)} rows, variations {counts}")
    print(f"  {OUT_GOLD.name}: gold tuples written")


if __name__ == "__main__":
    build()
