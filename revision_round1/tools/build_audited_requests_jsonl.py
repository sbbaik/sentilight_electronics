# -*- coding: utf-8 -*-
"""Turn the author-audited workbook into the evaluation row format the hold-out pipeline consumes.

Rows mirror `holdout_2000.jsonl` field for field — same instruction string, absent attributes as null, gold
tuple as the `output` dict — so the unchanged evaluators score this set exactly as they scored the others.
The provenance that differs from the other sets (LLM draft, author audit, correction notes) is carried in
extra fields and in `audit_notes.md`, never silently.
"""
from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parents[1]
XLSX = HERE / "data/audited_requests/author_audit_K240.xlsx"
GOLD = HERE / "data/audited_requests/audited_cells_gold.json"
HOLDOUT = HERE / "data/datasets/holdout_natural_K240/holdout_2000.jsonl"
OUT = HERE / "data/audited_requests/audited_requests_K240.jsonl"
NOTES = HERE / "data/audited_requests/audit_notes.md"
PROV = HERE / "data/audited_requests/provenance.json"

DRAFT_MODEL = "chatGPT GPT-6 Astra (Extra High)"
DRAFT_DATE = "2026-10-21"
VARIATION_EN = {"일반": "plain", "부정형": "negation", "어순변형": "word_order", "맥락추가": "context_clause"}


def none(v):
    return None if v in (None, "none", "") else v


def main() -> None:
    instruction = json.loads(HOLDOUT.open(encoding="utf-8").readline())["instruction"]
    gold = {r["cell_id"]: r for r in json.loads(GOLD.read_text(encoding="utf-8"))["rows"]}
    ws = load_workbook(XLSX, data_only=True)["문장작성"]

    rows, notes = [], []
    for r in range(2, ws.max_row + 1):
        no, cell_id, _, _, _, variation, _, text, memo = (ws.cell(row=r, column=c).value for c in range(1, 10))
        if cell_id is None:
            continue
        text = str(text or "").strip()
        if not text:
            sys.exit(f"row {no} ({cell_id}) has no sentence")
        g = gold[cell_id]
        if g["row"] != no or g["variation"] != variation:
            sys.exit(f"row {no} ({cell_id}) disagrees with the generated gold")
        rows.append({
            "instruction": instruction,
            "input": text,
            "output": {k: g["tup"][i] for i, k in enumerate(("H", "S", "B", "Dimmer", "CT"))},
            "emotion": none(g["emotion"]),
            "base_color": none(g["colour"]),
            "intensity": g["intensity"],
            "source": f"audited_request_{VARIATION_EN[variation]}",
            "cell_id": cell_id,
            "variation": VARIATION_EN[variation],
            "audit_note": str(memo).strip() if memo else None,
        })
        if memo:
            notes.append((no, cell_id, variation, text, str(memo).strip()))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")
    digest = hashlib.sha256(OUT.read_bytes()).hexdigest()

    NOTES.write_text(
        "# Audit notes on the audited request set\n\n"
        f"{len(notes)} of {len(rows)} rows carry a note from the first author's line-by-line audit of the "
        f"LLM draft. Notes are reproduced verbatim.\n\n"
        "| row | cell | variation | sentence | note |\n|---|---|---|---|---|\n"
        + "\n".join(f"| {n} | {c} | {VARIATION_EN[v]} | {t} | {m} |" for n, c, v, t, m in notes) + "\n",
        encoding="utf-8")

    PROV.write_text(json.dumps({
        "set": "audited request set (K240)",
        "rows": len(rows),
        "sha256": digest,
        "draft": {"model": DRAFT_MODEL, "called_on": DRAFT_DATE,
                  "input_given": "the 240 policy cells (colour, emotion, intensity) and a variation type per cell; "
                                 "no training template or training row"},
        "audit": {"by": "first author", "scope": "every row read; sentences corrected where the draft did not "
                                                 "express the cell or read unnaturally; notes left on judgement calls",
                  "templates_consulted": False, "rows_with_notes": len(notes)},
        "gold": "computed by the executable policy from the cell; the writer labels nothing",
        "variation_counts": {v: sum(1 for x in rows if x["variation"] == v) for v in VARIATION_EN.values()},
        "credentials_in_release": False,
    }, indent=1, ensure_ascii=False), encoding="utf-8")

    print(f"  {OUT.name}: {len(rows)} rows, sha256 {digest[:16]}")
    print(f"  {NOTES.name}: {len(notes)} notes | {PROV.name} written")


if __name__ == "__main__":
    main()
